"""Snapshot the office Perc Test workbook into the generator's template snapshot.

Usage: python tools/snapshot_perc_test.py <Perc Test.xlsx> <out_dir>

Cells are located by their labels, so a re-laid-out template still snapshots. Input
cells (header values, reading rows, tester name and date) are cleared, and the values
embedded in label text ("Test hole dia: 6\"") are cut back to the label, so sample job
data never enters the snapshot.
"""
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from snapshot_soil_log import style_of  # noqa: E402

# label → input key; the value cell is the first cell right of the label's merged range.
VALUE_LABELS = {
    "owner name:": "ownerName",
    "project name:": "projectName",
    "began:": "soakBegan",
    "begin:": "soakBegan",
    "ended:": "soakEnded",
    "end:": "soakEnded",
    "test date:": "testDate",
    "confirmation number:": "confirmationNumber",
}
# labels that carry their value in the same cell
EMBEDDED_LABELS = {
    "test hole dia:": "holeDiameter",
    "reference point elevation above hole bottom:": "referenceHeight",
    "test #": "title",
}
TABLE_COLUMNS = {
    "start time": "start",
    "end time": "end",
    "time interval": "interval",
    "initial distance": "initial",
    "final distance": "final",
    "drop in": "drop",
    "percolation rate": "rate",
}
# caption under a line → key of the cell on the line above
CAPTIONS = {"name (printed)": "testerName", "signature": "signature", "date": "certDate", "company": "company"}


def norm(v):
    return re.sub(r"\s+", " ", v).strip().lower() if isinstance(v, str) else None


def main(src, out_dir):
    raw = open(src, "rb").read()
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    merged = {}
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng

    def right_of(c):
        rng = merged.get((c.row, c.column))
        return f"{get_column_letter((rng.max_col if rng else c.column) + 1)}{c.row}"

    inputs, labels, table, sig = {}, {}, None, {}
    for row in ws.iter_rows():
        for c in row:
            v = norm(c.value)
            if not v:
                continue
            if v in VALUE_LABELS:
                inputs[VALUE_LABELS[v]] = right_of(c)
            for stem, key in EMBEDDED_LABELS.items():
                if v.startswith(stem):
                    inputs[key] = c.coordinate
                    labels[key] = c.value[: len(stem)].rstrip(" #") if key != "title" else c.value.split("#")[0] + "#"
            if v.startswith("start time"):
                rng = merged.get((c.row, c.column))
                table = {"headerRow": c.row, "firstRow": (rng.max_row if rng else c.row) + 1, "columns": {}}
                for h in ws[c.row]:
                    for stem, key in TABLE_COLUMNS.items():
                        if norm(h.value) and norm(h.value).startswith(stem):
                            table["columns"][key] = h.column_letter
            if v in CAPTIONS and table and c.row > table["headerRow"]:
                sig[CAPTIONS[v]] = f"{c.column_letter}{c.row - 1}"
    missing = set(VALUE_LABELS.values()) - set(inputs) | set(EMBEDDED_LABELS.values()) - set(inputs)
    if missing or not table or len(table["columns"]) != len(TABLE_COLUMNS) or len(sig) != len(CAPTIONS):
        sys.exit(f"template labels not found: missing={missing} table={table} signature={sig}")

    # Reading rows: rows with a start time below the header.
    r, col = table["firstRow"], table["columns"]["start"]
    while ws[f"{col}{r}"].value not in (None, ""):
        r += 1
    table["rows"] = r - table["firstRow"]
    table["rowHeight"] = ws.row_dimensions[table["firstRow"]].height

    clear = {v for k, v in inputs.items() if k not in labels} | {sig["testerName"], sig["certDate"]}
    for i in range(table["rows"]):
        for c in table["columns"].values():
            clear.add(f"{c}{table['firstRow'] + i}")

    cells = {}
    for row in ws.iter_rows():
        for c in row:
            s = style_of(c)
            s.pop("fill", None)
            keep = c.value not in (None, "") and c.coordinate not in clear
            if not keep and "border" not in s and c.coordinate not in clear:
                continue
            if keep:
                key = next((k for k, a in inputs.items() if a == c.coordinate and k in labels), None)
                s["value"] = labels[key] if key else c.value
            cells[c.coordinate] = s

    ps = ws.page_setup
    m = ws.page_margins
    snapshot = {
        "kind": "perc-test",
        "source": {
            "path": r"Office\Tools\Wastewater Tools\SEPTIC\SITE EVALUATION\Perc Test.xlsx",
            "sha256": hashlib.sha256(raw).hexdigest(),
            "modified": datetime.fromtimestamp(os.path.getmtime(src), timezone.utc).isoformat(),
            "snapshotAt": datetime.now(timezone.utc).isoformat(),
        },
        "sheetName": ws.title,
        "printArea": str(ws.print_area).split("!")[-1].replace("$", "") if ws.print_area else None,
        "columns": {k: round(v.width, 4) for k, v in ws.column_dimensions.items() if v.width},
        "defaultColumnWidth": ws.sheet_format.defaultColWidth or 9.140625,
        "rows": {str(k): v.height for k, v in sorted(ws.row_dimensions.items()) if v.height},
        "merges": sorted(str(rng) for rng in ws.merged_cells.ranges),
        "pageSetup": {
            "paperSize": int(ps.paperSize or 1),
            "orientation": ps.orientation or "portrait",
            "fitToPage": bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage),
            "fitToWidth": 1,
            "fitToHeight": 1,
            "margins": {"left": m.left, "right": m.right, "top": m.top, "bottom": m.bottom,
                        "header": m.header, "footer": m.footer},
        },
        "inputs": inputs,
        "labels": labels,
        "table": table,
        "signature": sig,
        "cells": cells,
    }
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "snapshot.json"), "w", encoding="utf-8") as f:
        json.dump(snapshot, f, indent=1, default=str)
        f.write("\n")
    print(f"snapshot written: {len(cells)} cells, inputs {inputs}, table {table}, signature {sig}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
