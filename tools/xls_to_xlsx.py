"""Convert the office Soil Log Template.xls to .xlsx without Excel, for snapshot_soil_log.py.

Usage: python tools/xls_to_xlsx.py <source.xls> <out.xlsx> <logo image> <logo anchor json>

xlrd reads values and cell formats (fonts, borders, fills, alignment, sizes, notes); it cannot
read pictures, so the logo is re-inserted from the committed template folder at its committed
anchor. Only the first worksheet is converted.
"""
import json
import sys

import openpyxl
import xlrd
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# BIFF line style codes -> openpyxl border styles.
LINE = {1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair", 8: "mediumDashed"}
HORIZONTAL = {1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous"}
VERTICAL = {0: "top", 1: "center", 2: "bottom", 3: "justify"}


def rgb(book, index):
    c = book.colour_map.get(index)
    return None if c is None else "FF%02X%02X%02X" % c


def convert(src, out, logo_path, anchor):
    book = xlrd.open_workbook(src, formatting_info=True)
    sh = book.sheet_by_index(0)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sh.name

    for c, info in sh.colinfo_map.items():
        if c < 26:
            ws.column_dimensions[get_column_letter(c + 1)].width = info.width / 256
    for r, info in sh.rowinfo_map.items():
        if not info.height_mismatch and info.height == 255:
            continue
        ws.row_dimensions[r + 1].height = info.height / 20

    for r in range(sh.nrows):
        for c in range(sh.ncols):
            xf = book.xf_list[sh.cell_xf_index(r, c)]
            cell = ws.cell(r + 1, c + 1)
            v = sh.cell_value(r, c)
            if v != "":
                cell.value = v
            f = book.font_list[xf.font_index]
            cell.font = Font(name=f.name, size=f.height / 20, bold=bool(f.bold), italic=bool(f.italic))
            fmt = book.format_map.get(xf.format_key)
            if fmt:
                cell.number_format = fmt.format_str
            a = xf.alignment
            cell.alignment = Alignment(horizontal=HORIZONTAL.get(a.hor_align), vertical=VERTICAL.get(a.vert_align), wrap_text=bool(a.text_wrapped))
            b = xf.border
            side = lambda style, colour: Side(style=LINE.get(style)) if style else Side()
            cell.border = Border(left=side(b.left_line_style, b.left_colour_index), right=side(b.right_line_style, b.right_colour_index),
                                 top=side(b.top_line_style, b.top_colour_index), bottom=side(b.bottom_line_style, b.bottom_colour_index))
            bg = xf.background
            if bg.fill_pattern == 1:
                colour = rgb(book, bg.pattern_colour_index)
                if colour and colour != "FFFFFFFF":
                    cell.fill = PatternFill("solid", fgColor=colour)
    for (r, c), note in sh.cell_note_map.items():
        ws.cell(r + 1, c + 1).comment = Comment(note.text, note.author or "")

    ws.page_setup.paperSize = 1
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    img = Image(logo_path)
    emu = 9525
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=anchor["from"]["col"], colOff=anchor["from"]["colOff"], row=anchor["from"]["row"], rowOff=anchor["from"]["rowOff"]),
        ext=XDRPositiveSize2D(anchor["extEmu"]["cx"], anchor["extEmu"]["cy"]),
    )
    img.width, img.height = anchor["extEmu"]["cx"] / emu, anchor["extEmu"]["cy"] / emu
    ws.add_image(img)
    wb.save(out)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2], sys.argv[3], json.loads(open(sys.argv[4], encoding="utf-8-sig").read()))
